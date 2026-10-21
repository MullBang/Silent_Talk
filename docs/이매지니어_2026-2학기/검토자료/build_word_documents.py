from pathlib import Path
import re
import json
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
FONT = '맑은 고딕'

def font_spec(style, size, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for a in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        rf.set(qn('w:' + a), FONT)
    for a in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        rf.attrib.pop(qn('w:' + a), None)

def setup(title, subtitle):
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin, sec.bottom_margin = Mm(15), Mm(15)
    sec.left_margin, sec.right_margin = Mm(20), Mm(20)
    sec.header_distance = sec.footer_distance = Mm(10)
    for name, size, bold in [('Normal',12,False),('Title',20,True),('Subtitle',12,False),('Heading 1',15,True),('Heading 2',13,True),('Heading 3',12,True),('Heading 4',12,True),('List Bullet',12,False),('Caption',10,False)]:
        s = d.styles[name]
        font_spec(s,size,bold)
        p = s.paragraph_format
        p.line_spacing = 1.6
        p.space_after = Pt(5)
        p.widow_control = True
        if name.startswith('Heading'):
            p.space_before = Pt(12 if name=='Heading 1' else 8)
            p.space_after = Pt(4)
            p.keep_with_next = True
        else:
            p.keep_with_next = False
        if name in ['Title','Subtitle']:
            p.line_spacing=1.15
            p.keep_with_next=True
            p.space_after=Pt(7)
        pp=s.element.find(qn('w:pPr'))
        if pp is not None:
            for tag in ['pBdr','shd']:
                for e in pp.findall(qn('w:'+tag)):pp.remove(e)
    d.styles['Normal'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    d.styles['List Bullet'].paragraph_format.left_indent=Mm(4)
    d.styles['List Bullet'].paragraph_format.first_line_indent=Mm(-3)
    for s in d.styles:
        if s.type==1:
            s.font.color.rgb=RGBColor(0,0,0)
    settings=d.settings.element
    lang=OxmlElement('w:themeFontLang')
    lang.set(qn('w:val'),'en-US');lang.set(qn('w:eastAsia'),'ko-KR')
    settings.append(lang)
    d.core_properties.title=title
    d.core_properties.subject=subtitle
    d.core_properties.author=''
    d.core_properties.last_modified_by=''
    d.core_properties.comments=''
    d.add_paragraph(title, 'Title')
    d.add_paragraph(subtitle, 'Subtitle')
    f=sec.footer.paragraphs[0]
    f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    f.paragraph_format.space_before=Pt(0);f.paragraph_format.space_after=Pt(0)
    r=f.add_run();r.font.size=Pt(9);r.font.name=FONT
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
    rr=OxmlElement('w:r');tt=OxmlElement('w:t');tt.text='1';rr.append(tt);field.append(rr);f._p.append(field)
    return d

def heading_text(s):
    s=s.strip('[] ')
    s=re.sub(r'^○\s*','',s)
    s=re.sub(r'^[①②③④⑤⑥⑦⑧⑨⑩]\s*','',s)
    s=re.sub(r'^(\d+)\.\s*',r'\1 ',s)
    s=re.sub(r'^(\d+)\)\s*',r'\1 ',s)
    s=re.sub(r'^([가나다라마바사아])\.\s*',r'\1 ',s)
    s=s.replace('2026-1학기','2026년 1학기').replace('2026-2학기','2026년 2학기')
    s=re.sub(r'[—·/()]',' ',s)
    s=s.replace('‘','').replace('’','')
    return re.sub(r'\s+',' ',s).strip()

def para(d,text,style=None):
    p=d.add_paragraph(style=style)
    if text.startswith('※ '):
        p.add_run(text[2:])
        p.paragraph_format.space_before=Pt(3)
    elif ': ' in text and not text.startswith(('http','C:')) and len(text.split(': ',1)[0])<35:
        a,b=text.split(': ',1)
        p.add_run(a+': ').bold=True;p.add_run(b)
    else:p.add_run(text)
    return p

def table(d,rows,widths,centers=()):
    t=d.add_table(rows=1,cols=len(widths))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    pr=t._tbl.tblPr
    bd=OxmlElement('w:tblBorders')
    for k in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+k);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');bd.append(e)
    pr.append(bd)
    mar=OxmlElement('w:tblCellMar')
    for k,v in [('top',80),('left',100),('bottom',80),('right',100)]:
        e=OxmlElement('w:'+k);e.set(qn('w:w'),str(v));e.set(qn('w:type'),'dxa');mar.append(e)
    pr.append(mar)
    for c,w in zip(t.columns,widths):c.width=Mm(w)
    for ri,values in enumerate(rows):
        row=t.rows[0] if ri==0 else t.add_row()
        trp=row._tr.get_or_add_trPr()
        ns=OxmlElement('w:cantSplit');trp.append(ns)
        if ri==0:
            rh=OxmlElement('w:tblHeader');trp.append(rh)
        for ci,(c,txt,w) in enumerate(zip(row.cells,values,widths)):
            c.width=Mm(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E5EDF3' if ri==0 else 'FFFFFF');c._tc.get_or_add_tcPr().append(sh)
            p=c.paragraphs[0]
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ri==0 or ci in centers else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing=1.2
            p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.keep_with_next=False
            r=p.add_run(txt);r.font.name=FONT;r.font.size=Pt(10.5);r.bold=(ri==0)
            r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),FONT)
    p=d.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
    p.add_run().font.size=Pt(3)
    return t

def prepare_lines(path,kind):
    ls=path.read_text(encoding='utf-8-sig').splitlines()[1:]
    if kind=='01':
        a=ls.index('[사용 안내]');b=ls.index('1. 주제 선정 배경 및 필요성')
        guide=[s for s in ls[a+1:b] if s.strip()]
        ls=ls[:a]+ls[b:]+['','[사용 안내]']+guide
    result=[]
    for s in ls:
        for stem in ['01_서식기입용_작성안','02_작성근거와_확인사항','03_핵심항목_요약본']:
            s=s.replace(stem+'.txt',stem+'.docx')
        s=s.replace('TXT에는 페이지 개념이 없어 실제 HWP 배치 후 분량을 확인해야 한다.','상세 작성안에는 선택용 문구와 참고 안내가 포함되어 있으므로, 제출 시 필요한 내용을 선별하여 실제 서식에 맞춰 분량을 조정해야 한다.')
        result.append(s)
    return result

def build(stem,title,subtitle):
    kind=stem[:2]
    d=setup(title,subtitle)
    ls=prepare_lines(ROOT/(stem+'.txt'),kind)
    i=0;h1='';all_source=[]
    while i<len(ls):
        s=ls[i].strip()
        if not s:i+=1;continue
        if ' | ' in s:
            rs=[]
            while i<len(ls) and ' | ' in ls[i]:
                rs.append([x.strip() for x in ls[i].split(' | ')]);i+=1
            widths=[32,35,103] if len(rs[0])==3 else [54,22,32,25,37]
            table(d,rs,widths,(1,) if len(widths)==3 else (1,2,3,4));continue
        if kind=='01' and s=='○ 추진일정':
            d.add_paragraph('추진일정','Heading 2');i+=1
            rows=[['운영기간','추진내용']]
            while i<len(ls):
                if not ls[i].strip():i+=1;continue
                if not re.match(r'^2026\.\d{2}\.\d{2}',ls[i]):break
                date=ls[i].strip();i+=1
                while not ls[i].strip():i+=1
                rows.append([date,ls[i].strip()]);i+=1
            table(d,rows,[39,131],(0,));continue
        if kind=='01' and s.startswith('지도교수:'):
            rows=[['참여자','역할 및 활동내용']]
            while i<len(ls) and any(ls[i].startswith(x+':') for x in ['지도교수','참여학생','참여기업']):
                a,b=ls[i].split(':',1);rows.append([a,b.strip()]);i+=1
            table(d,rows,[26,144],(0,));continue
        if re.match(r'^\d+\. ',s):
            h1=s
            p=d.add_paragraph(heading_text(s),'Heading 1')
            if kind=='01' and s.startswith(('2. ','3. ','4. ')):p.paragraph_format.page_break_before=True
            if kind=='02' and s.startswith(('2. ','3. ','4. ')):p.paragraph_format.page_break_before=True
        elif s.startswith('○ ') or (s.startswith('[') and s.endswith(']')):
            p=d.add_paragraph(heading_text(s),'Heading 2')
            if kind=='01' and (s.startswith('[2026-2학기') or s.startswith('○ 정량적') or s.startswith('○ 필수활동')):
                p.paragraph_format.page_break_before=True
        elif re.match(r'^[①②③④⑤⑥⑦⑧⑨⑩]',s) or re.match(r'^\d+\) ',s) or re.match(r'^[가나다라마바사아]\. ',s):
            d.add_paragraph(heading_text(s),'Heading 3')
        elif s.startswith('- '):para(d,s[2:],'List Bullet')
        elif s.startswith('참여기업(기관) 확인서 —'):
            d.add_paragraph('참여기업 확인서 주제 및 내용','Heading 3')
        else:para(d,s)
        i+=1
    out=ROOT/(stem+'.docx')
    d.save(out)
    check=Document(out)
    stats={'filename':out.name,'paragraphs':len(check.paragraphs),'tables':len(check.tables),'bytes':out.stat().st_size}
    return stats

if __name__=='__main__':
    stats=[]
    stats.append(build('01_서식기입용_작성안','2026년 2학기 이매지니어 프로젝트 운영계획서','서식 기입용 상세 작성안'))
    stats.append(build('02_작성근거와_확인사항','이매지니어 프로젝트 작성 근거와 확인사항','운영계획서 작성 참고자료'))
    stats.append(build('03_핵심항목_요약본','2026년 2학기 이매지니어 프로젝트 핵심항목 요약','배경과 필요성  진행 성과  운영 목표  기대효과'))
    (ROOT/'검토자료'/'word_build_report.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(stats,ensure_ascii=True))
